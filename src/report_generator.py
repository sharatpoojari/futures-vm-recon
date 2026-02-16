"""Report Generator Module for creating Excel reconciliation reports"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime
from typing import Dict
from pathlib import Path


class ReportGenerator:
    """Generates formatted Excel reports for reconciliation results"""

    def __init__(self):
        """Initialize ReportGenerator"""
        self.red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        self.yellow_fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
        self.green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
        self.light_blue_fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF")

    def generate_excel_report(self, recon_results: Dict[str, pd.DataFrame], output_path: str) -> str:
        """
        Generate comprehensive Excel reconciliation report

        Args:
            recon_results: Dictionary containing reconciliation results
            output_path: Path where the Excel report should be saved

        Returns:
            Path to the generated report
        """
        # Create workbook
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Create sheets
        self._create_summary_sheet(wb, recon_results)
        self._create_matched_sheet(wb, recon_results['matched'])
        self._create_exceptions_sheet(wb, recon_results['fund_admin_only'], 'Fund Admin Only')
        self._create_exceptions_sheet(wb, recon_results['custody_only'], 'Custody Only')
        self._create_mismatches_sheet(wb, recon_results['amount_mismatches'])

        # Save workbook
        wb.save(output_path)
        return output_path

    def _create_summary_sheet(self, wb: Workbook, recon_results: Dict[str, pd.DataFrame]) -> None:
        """Create summary sheet with high-level statistics"""
        ws = wb.create_sheet("Summary", 0)

        # Calculate statistics
        total_fa = len(recon_results['matched']) + len(recon_results['fund_admin_only']) + len(recon_results['amount_mismatches'])
        total_cust = len(recon_results['matched']) + len(recon_results['custody_only']) + len(recon_results['amount_mismatches'])
        matched_count = len(recon_results['matched'])
        fa_only_count = len(recon_results['fund_admin_only'])
        cust_only_count = len(recon_results['custody_only'])
        mismatch_count = len(recon_results['amount_mismatches'])

        match_rate = (matched_count / total_fa * 100) if total_fa > 0 else 0

        # Header
        ws['A1'] = 'Futures Variation Margin Reconciliation Report'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'

        # Statistics table
        row = 4
        ws[f'A{row}'] = 'Metric'
        ws[f'B{row}'] = 'Value'
        self._apply_header_style(ws, row, 2)

        row += 1
        ws[f'A{row}'] = 'Total Fund Admin Records'
        ws[f'B{row}'] = total_fa

        row += 1
        ws[f'A{row}'] = 'Total Custody Records'
        ws[f'B{row}'] = total_cust

        row += 1
        ws[f'A{row}'] = 'Matched Records'
        ws[f'B{row}'] = matched_count
        ws[f'B{row}'].fill = self.green_fill

        row += 1
        ws[f'A{row}'] = 'Fund Admin Only'
        ws[f'B{row}'] = fa_only_count
        if fa_only_count > 0:
            ws[f'B{row}'].fill = self.red_fill

        row += 1
        ws[f'A{row}'] = 'Custody Only'
        ws[f'B{row}'] = cust_only_count
        if cust_only_count > 0:
            ws[f'B{row}'].fill = self.red_fill

        row += 1
        ws[f'A{row}'] = 'Amount Mismatches (Total)'
        ws[f'B{row}'] = mismatch_count
        if mismatch_count > 0:
            ws[f'B{row}'].fill = self.yellow_fill

        # Fee adjustment breakdown
        mismatches_df = recon_results['amount_mismatches']
        if not mismatches_df.empty and 'status' in mismatches_df.columns:
            explained_count = len(mismatches_df[mismatches_df['status'] == 'Explained by Fees'])
            unexplained_count = len(mismatches_df[mismatches_df['status'] != 'Explained by Fees'])
        else:
            explained_count = 0
            unexplained_count = mismatch_count

        row += 1
        ws[f'A{row}'] = '  - Explained by Fees'
        ws[f'B{row}'] = explained_count
        if explained_count > 0:
            ws[f'B{row}'].fill = self.light_blue_fill

        row += 1
        ws[f'A{row}'] = '  - Unexplained Mismatches'
        ws[f'B{row}'] = unexplained_count
        if unexplained_count > 0:
            ws[f'B{row}'].fill = self.yellow_fill

        row += 2
        ws[f'A{row}'] = 'Match Rate'
        ws[f'B{row}'] = f'{match_rate:.2f}%'
        ws[f'B{row}'].font = Font(bold=True)

        # Adjust column widths
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 20

    def _create_matched_sheet(self, wb: Workbook, matched_df: pd.DataFrame) -> None:
        """Create sheet for matched records"""
        ws = wb.create_sheet("Matched")

        if matched_df.empty:
            ws['A1'] = 'No matched records'
            return

        # Write data
        for r_idx, row in enumerate(dataframe_to_rows(matched_df, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)

                # Format header row
                if r_idx == 1:
                    self._apply_header_style(ws, r_idx, len(row))

        # Auto-adjust column widths
        self._auto_adjust_columns(ws, matched_df)

    def _create_exceptions_sheet(self, wb: Workbook, exceptions_df: pd.DataFrame, sheet_name: str) -> None:
        """Create sheet for exception records (fund admin only or custody only)"""
        ws = wb.create_sheet(sheet_name)

        if exceptions_df.empty:
            ws['A1'] = f'No {sheet_name.lower()} records'
            return

        # Write data
        for r_idx, row in enumerate(dataframe_to_rows(exceptions_df, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)

                # Format header row
                if r_idx == 1:
                    self._apply_header_style(ws, r_idx, len(row))
                # Highlight exception rows
                elif r_idx > 1:
                    cell.fill = self.red_fill

        # Auto-adjust column widths
        self._auto_adjust_columns(ws, exceptions_df)

    def _create_mismatches_sheet(self, wb: Workbook, mismatches_df: pd.DataFrame) -> None:
        """Create sheet for amount mismatches with conditional coloring"""
        ws = wb.create_sheet("Amount Mismatches")

        if mismatches_df.empty:
            ws['A1'] = 'No amount mismatches'
            return

        has_status = 'status' in mismatches_df.columns

        # Write data
        for r_idx, row in enumerate(dataframe_to_rows(mismatches_df, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)

                # Format header row
                if r_idx == 1:
                    self._apply_header_style(ws, r_idx, len(row))
                # Conditional row coloring based on status
                elif r_idx > 1:
                    if has_status:
                        row_status = mismatches_df.iloc[r_idx - 2]['status']
                        if row_status == 'Explained by Fees':
                            cell.fill = self.light_blue_fill
                        else:
                            cell.fill = self.yellow_fill
                    else:
                        cell.fill = self.yellow_fill

        # Auto-adjust column widths
        self._auto_adjust_columns(ws, mismatches_df)

    def _apply_header_style(self, ws, row: int, col_count: int) -> None:
        """Apply header styling to a row"""
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

    def _auto_adjust_columns(self, ws, df: pd.DataFrame) -> None:
        """Auto-adjust column widths based on content"""
        for idx, col in enumerate(df.columns, 1):
            max_length = max(
                df[col].astype(str).apply(len).max(),
                len(str(col))
            )
            adjusted_width = min(max_length + 2, 50)  # Cap at 50
            ws.column_dimensions[get_column_letter(idx)].width = adjusted_width
